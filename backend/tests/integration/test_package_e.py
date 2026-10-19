"""Integration tests for Package E: EaseGuide (assistance), EaseVoice (narration/settings), EaseAnalytics."""
from datetime import timedelta
from types import SimpleNamespace

import pytest
from alembic import command
from sqlalchemy import inspect, select

from app.core.config import get_settings
from app.core.security import utcnow
from app.db.models import AuditLog, Event, Role, TicketCategory
from app.db.session import engine
from app.services.payments.base import MockPaymentProvider, set_payment_provider
from tests.helpers import admin_headers, auth_header, error_of, make_event, make_user


@pytest.fixture(autouse=True)
def mock_provider():
    set_payment_provider(MockPaymentProvider())
    yield
    set_payment_provider(None)


def on_sale_event(db, organiser, quantity=50, price=5000, **fields):
    event = make_event(db, organiser=organiser, days=30, categories=[("STANDARD", price, quantity)], **fields)
    category = db.scalar(select(TicketCategory).where(TicketCategory.event_id == event.id))
    return event, category


def create_booking(client, headers, event, category, quantity=1):
    response = client.post("/api/v1/bookings", headers=headers,
                           json={"eventId": str(event.id),
                                 "items": [{"ticketCategoryId": str(category.id), "quantity": quantity}]})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def confirmed_booking(client, headers, event, category, quantity=1):
    booking_id = create_booking(client, headers, event, category, quantity)
    paid = client.post(f"/api/v1/bookings/{booking_id}/pay", headers=headers, json={})
    assert paid.status_code == 200, paid.text
    assert paid.json()["status"] == "CONFIRMED"
    return booking_id


def request_assistance(client, headers, booking_id, **overrides):
    body = {"bookingId": booking_id, "supportNeeds": "Wheelchair access", "assistanceLevel": "MODERATE", **overrides}
    return client.post("/api/v1/assistance-requests", headers=headers, json=body)


@pytest.fixture
def world(client, db):
    organiser = make_user(db, role=Role.ORGANISER)
    event, category = on_sale_event(db, organiser)
    customer = make_user(db)
    return SimpleNamespace(
        organiser=organiser, event=event, category=category, customer=customer,
        headers=auth_header(client, customer.username),
        organiser_headers=auth_header(client, organiser.username),
        admin=admin_headers(client, db),
    )


# ------------------------------------------------------------------ EaseGuide: ticket holders

def test_customer_requests_assistance_for_a_confirmed_booking(client, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    response = request_assistance(client, world.headers, booking_id)
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "SUBMITTED" and body["assistanceLevel"] == "MODERATE"
    assert body["assignedStaff"] is None and body["bookingId"] == booking_id
    assert body["event"]["title"] == world.event.title
    assert "userId" not in body  # staff-only field

    listing = client.get("/api/v1/me/assistance-requests", headers=world.headers).json()
    assert listing["total"] == 1 and listing["items"][0]["id"] == body["id"]


def test_assistance_needs_a_confirmed_booking(client, world):
    unpaid = create_booking(client, world.headers, world.event, world.category)
    response = request_assistance(client, world.headers, unpaid)
    assert response.status_code == 409
    assert error_of(response)["code"] == "BOOKING_NOT_CONFIRMED"


def test_cannot_request_assistance_for_someone_elses_booking(client, db, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    stranger = auth_header(client, make_user(db).username)
    assert request_assistance(client, stranger, booking_id).status_code == 404


def test_only_one_open_request_per_booking(client, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    first = request_assistance(client, world.headers, booking_id)
    assert first.status_code == 201
    duplicate = request_assistance(client, world.headers, booking_id)
    assert duplicate.status_code == 409
    error = error_of(duplicate)
    assert error["code"] == "ASSISTANCE_ALREADY_REQUESTED"
    assert error["details"]["requestId"] == first.json()["id"]


def test_a_new_request_is_allowed_after_the_previous_one_is_cancelled(client, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    first = request_assistance(client, world.headers, booking_id).json()
    cancelled = client.post(f"/api/v1/me/assistance-requests/{first['id']}/cancel", headers=world.headers)
    assert cancelled.status_code == 200 and cancelled.json()["status"] == "CANCELLED"
    assert request_assistance(client, world.headers, booking_id).status_code == 201


def test_no_assistance_requests_once_the_event_has_ended(client, db, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    event = db.get(Event, world.event.id)
    event.start_at = utcnow() - timedelta(hours=3)
    event.end_at = utcnow() - timedelta(hours=1)
    db.commit()
    response = request_assistance(client, world.headers, booking_id)
    assert response.status_code == 409
    assert error_of(response)["code"] == "EVENT_ALREADY_ENDED"


def test_request_validation(client, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    assert request_assistance(client, world.headers, booking_id, supportNeeds="  ").status_code == 422
    assert request_assistance(client, world.headers, booking_id, supportNeeds="x" * 1001).status_code == 422
    assert request_assistance(client, world.headers, booking_id, assistanceLevel="EXTREME").status_code == 422


def test_customers_only_see_their_own_requests(client, db, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    request_id = request_assistance(client, world.headers, booking_id).json()["id"]
    other = auth_header(client, make_user(db).username)
    assert client.get(f"/api/v1/me/assistance-requests/{request_id}", headers=other).status_code == 404
    assert client.get("/api/v1/me/assistance-requests", headers=other).json()["total"] == 0
    assert client.post(f"/api/v1/me/assistance-requests/{request_id}/cancel", headers=other).status_code == 404
    bad_id = client.get("/api/v1/me/assistance-requests/not-a-uuid", headers=world.headers)
    assert bad_id.status_code == 404 and error_of(bad_id)["code"] == "ASSISTANCE_REQUEST_NOT_FOUND"


def test_a_completed_request_cannot_be_cancelled_by_the_customer(client, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    request_id = request_assistance(client, world.headers, booking_id).json()["id"]
    admin_url = f"/api/v1/admin/assistance-requests/{request_id}"
    client.patch(admin_url, headers=world.admin, json={"assignedStaff": "Alex Tan"})
    client.patch(admin_url, headers=world.admin, json={"status": "IN_PROGRESS"})
    client.patch(admin_url, headers=world.admin, json={"status": "COMPLETED"})
    response = client.post(f"/api/v1/me/assistance-requests/{request_id}/cancel", headers=world.headers)
    assert response.status_code == 409
    assert error_of(response)["code"] == "ASSISTANCE_NOT_CANCELLABLE"


def test_customer_endpoints_are_for_customers_only(client, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    body = {"bookingId": booking_id, "supportNeeds": "Wheelchair access", "assistanceLevel": "BASIC"}
    assert client.post("/api/v1/assistance-requests", json=body).status_code == 401
    for headers in (world.admin, world.organiser_headers):
        response = client.post("/api/v1/assistance-requests", headers=headers, json=body)
        assert response.status_code == 403 and error_of(response)["code"] == "FORBIDDEN"
        assert client.get("/api/v1/me/assistance-requests", headers=headers).status_code == 403


# ------------------------------------------------------------------ EaseGuide: venue staff

def test_staff_lifecycle_is_audited_end_to_end(client, db, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    request_id = request_assistance(client, world.headers, booking_id).json()["id"]
    url = f"/api/v1/admin/assistance-requests/{request_id}"

    assigned = client.patch(url, headers=world.admin, json={"assignedStaff": "Alex Tan"})
    assert assigned.status_code == 200, assigned.text
    assert assigned.json()["status"] == "ASSIGNED" and assigned.json()["assignedStaff"] == "Alex Tan"
    assert assigned.json()["requesterUsername"] == world.customer.username

    for target in ("IN_PROGRESS", "COMPLETED"):
        step = client.patch(url, headers=world.admin, json={"status": target})
        assert step.status_code == 200, step.text
        assert step.json()["status"] == target

    mine = client.get(f"/api/v1/me/assistance-requests/{request_id}", headers=world.headers).json()
    assert mine["status"] == "COMPLETED" and mine["assignedStaff"] == "Alex Tan"

    db.expire_all()
    rows = db.scalars(select(AuditLog).where(AuditLog.entity_type == "AssistanceRequest")
                      .order_by(AuditLog.created_at)).all()
    assert [r.action for r in rows] == ["ASSISTANCE_REQUESTED"] + ["ASSISTANCE_UPDATED"] * 3
    assert [r.actor_role for r in rows] == ["CUSTOMER", "ADMIN", "ADMIN", "ADMIN"]
    assert [r.after["status"] for r in rows] == ["SUBMITTED", "ASSIGNED", "IN_PROGRESS", "COMPLETED"]
    assert rows[1].before["status"] == "SUBMITTED"


def test_staff_cannot_skip_statuses_or_reopen_closed_requests(client, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    request_id = request_assistance(client, world.headers, booking_id).json()["id"]
    url = f"/api/v1/admin/assistance-requests/{request_id}"

    skipped = client.patch(url, headers=world.admin, json={"status": "COMPLETED"})
    assert skipped.status_code == 409
    error = error_of(skipped)
    assert error["code"] == "INVALID_STATUS_TRANSITION"
    assert error["details"] == {"from": "SUBMITTED", "to": "COMPLETED"}

    no_staff = client.patch(url, headers=world.admin, json={"status": "ASSIGNED"})
    assert no_staff.status_code == 400 and error_of(no_staff)["code"] == "ASSIGNED_STAFF_REQUIRED"
    assert client.patch(url, headers=world.admin, json={}).status_code == 422
    assert client.patch(url, headers=world.admin, json={"assignedStaff": "   "}).status_code == 422

    assert client.patch(url, headers=world.admin, json={"status": "CANCELLED"}).status_code == 200
    closed = client.patch(url, headers=world.admin, json={"status": "IN_PROGRESS"})
    assert closed.status_code == 409 and error_of(closed)["code"] == "ASSISTANCE_REQUEST_CLOSED"


def test_staff_can_reassign_an_open_request(client, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    request_id = request_assistance(client, world.headers, booking_id).json()["id"]
    url = f"/api/v1/admin/assistance-requests/{request_id}"
    client.patch(url, headers=world.admin, json={"assignedStaff": "Alex Tan"})
    reassigned = client.patch(url, headers=world.admin, json={"assignedStaff": "Sam Lee"})
    assert reassigned.status_code == 200
    assert reassigned.json()["status"] == "ASSIGNED" and reassigned.json()["assignedStaff"] == "Sam Lee"


def test_staff_queue_filters_and_paginates(client, world):
    ids = []
    for _ in range(2):
        booking_id = confirmed_booking(client, world.headers, world.event, world.category)
        ids.append(request_assistance(client, world.headers, booking_id).json()["id"])
    client.patch(f"/api/v1/admin/assistance-requests/{ids[0]}", headers=world.admin,
                 json={"assignedStaff": "Alex Tan"})

    everything = client.get("/api/v1/admin/assistance-requests", headers=world.admin).json()
    assert everything["total"] == 2
    assert {i["requesterUsername"] for i in everything["items"]} == {world.customer.username}
    assert client.get("/api/v1/admin/assistance-requests?status=ASSIGNED", headers=world.admin).json()["total"] == 1
    assert client.get("/api/v1/admin/assistance-requests?status=SUBMITTED", headers=world.admin).json()["total"] == 1
    paged = client.get("/api/v1/admin/assistance-requests?pageSize=1", headers=world.admin).json()
    assert len(paged["items"]) == 1 and paged["totalPages"] == 2


def test_staff_endpoints_are_for_admins_only(client, world):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    request_id = request_assistance(client, world.headers, booking_id).json()["id"]
    url = f"/api/v1/admin/assistance-requests/{request_id}"
    for headers in (world.headers, world.organiser_headers):
        assert client.get("/api/v1/admin/assistance-requests", headers=headers).status_code == 403
        assert client.patch(url, headers=headers, json={"assignedStaff": "Alex Tan"}).status_code == 403
    assert client.get("/api/v1/admin/assistance-requests").status_code == 401
    unknown = client.patch("/api/v1/admin/assistance-requests/00000000-0000-0000-0000-000000000000",
                           headers=world.admin, json={"assignedStaff": "Alex Tan"})
    assert unknown.status_code == 404


# ------------------------------------------------------------------ EaseVoice

def test_narration_is_public_to_read_and_admin_only_to_write(client, world):
    # clean_tables empties the seed rows, so create our own.
    assert client.get("/api/v1/narration/home").status_code == 404
    created = client.put("/api/v1/admin/narration/home", headers=world.admin, json={"transcript": "  Welcome  "})
    assert created.status_code == 200, created.text
    assert created.json()["pageKey"] == "home" and created.json()["transcript"] == "Welcome"

    public = client.get("/api/v1/narration/home")  # no login
    assert public.status_code == 200 and public.json()["transcript"] == "Welcome"

    client.put("/api/v1/admin/narration/home", headers=world.admin, json={"transcript": "Updated"})
    assert client.get("/api/v1/narration/home").json()["transcript"] == "Updated"

    missing = client.get("/api/v1/narration/nothing-here")
    assert missing.status_code == 404 and error_of(missing)["code"] == "NARRATION_NOT_FOUND"
    assert client.get("/api/v1/narration/Bad_Key").status_code == 422
    for headers in (world.headers, world.organiser_headers):
        assert client.put("/api/v1/admin/narration/home", headers=headers,
                          json={"transcript": "Hijack"}).status_code == 403
    assert client.put("/api/v1/admin/narration/home", json={"transcript": "x"}).status_code == 401
    assert client.put("/api/v1/admin/narration/home", headers=world.admin, json={"transcript": "   "}).status_code == 422


def test_accessibility_settings_default_then_partial_updates(client, db, world):
    assert client.get("/api/v1/me/accessibility-settings", headers=world.headers).json() == {
        "voiceSpeed": 1.0, "autoNarrate": False}

    speed = client.put("/api/v1/me/accessibility-settings", headers=world.headers, json={"voiceSpeed": 1.5})
    assert speed.status_code == 200 and speed.json() == {"voiceSpeed": 1.5, "autoNarrate": False}
    both = client.put("/api/v1/me/accessibility-settings", headers=world.headers, json={"autoNarrate": True})
    assert both.json() == {"voiceSpeed": 1.5, "autoNarrate": True}  # speed kept

    other = auth_header(client, make_user(db).username)
    assert client.get("/api/v1/me/accessibility-settings", headers=other).json()["voiceSpeed"] == 1.0


def test_accessibility_settings_validation_and_roles(client, world):
    for bad in (0.4, 2.1, 3):
        response = client.put("/api/v1/me/accessibility-settings", headers=world.headers, json={"voiceSpeed": bad})
        assert response.status_code == 422
    assert client.get("/api/v1/me/accessibility-settings").status_code == 401
    for headers in (world.admin, world.organiser_headers):  # every logged-in role can use EaseVoice
        assert client.get("/api/v1/me/accessibility-settings", headers=headers).status_code == 200


# ------------------------------------------------------------------ EaseAnalytics

def test_analytics_report_confirmed_sales_only(client, world):
    confirmed_booking(client, world.headers, world.event, world.category, quantity=2)
    create_booking(client, world.headers, world.event, world.category, quantity=1)  # unpaid: must not count

    sales = client.get("/api/v1/organiser/analytics/sales", headers=world.organiser_headers)
    assert sales.status_code == 200, sales.text
    sales = sales.json()
    assert sales["currency"] == "SGD"
    assert (sales["confirmedBookings"], sales["totalTicketsSold"], sales["totalRevenueCents"]) == (1, 2, 10000)
    assert [e["eventId"] for e in sales["events"]] == [str(world.event.id)]

    # Agrees with the existing per-event sales report.
    report = client.get(f"/api/v1/organiser/events/{world.event.id}/sales", headers=world.organiser_headers).json()
    assert report["ticketsSold"] == sales["totalTicketsSold"]
    assert report["grossRevenueCents"] == sales["totalRevenueCents"]

    velocity = client.get("/api/v1/organiser/analytics/revenue-velocity?days=7", headers=world.organiser_headers).json()
    assert velocity["days"] == 7 and len(velocity["points"]) == 7
    assert velocity["totalRevenueCents"] == 10000 and velocity["averageDailyRevenueCents"] == 10000 // 7
    today = velocity["points"][-1]
    assert (today["ticketsSold"], today["revenueCents"], today["cumulativeRevenueCents"]) == (2, 10000, 10000)
    assert all(p["revenueCents"] == 0 for p in velocity["points"][:-1])

    tiers = client.get("/api/v1/organiser/analytics/tier-demand", headers=world.organiser_headers).json()
    assert tiers["totalTicketsSold"] == 2 and len(tiers["tiers"]) == 4  # every tier is always returned
    by_type = {t["type"]: t for t in tiers["tiers"]}
    assert by_type["STANDARD"]["ticketsSold"] == 2 and by_type["STANDARD"]["sharePercent"] == 100.0
    assert by_type["VIP"]["ticketsSold"] == 0 and by_type["VIP"]["sharePercent"] == 0.0


def test_analytics_with_no_sales_are_zero_not_errors(client, world):
    create_booking(client, world.headers, world.event, world.category)  # held, never paid
    sales = client.get("/api/v1/organiser/analytics/sales", headers=world.organiser_headers).json()
    assert (sales["confirmedBookings"], sales["totalTicketsSold"], sales["totalRevenueCents"]) == (0, 0, 0)
    assert len(sales["events"]) == 1 and sales["events"][0]["ticketsSold"] == 0
    tiers = client.get("/api/v1/organiser/analytics/tier-demand", headers=world.organiser_headers).json()
    assert tiers["totalTicketsSold"] == 0 and all(t["sharePercent"] == 0.0 for t in tiers["tiers"])
    velocity = client.get("/api/v1/organiser/analytics/revenue-velocity", headers=world.organiser_headers).json()
    assert velocity["days"] == 30 and len(velocity["points"]) == 30 and velocity["totalRevenueCents"] == 0


def test_organisers_only_see_their_own_events(client, db, world):
    confirmed_booking(client, world.headers, world.event, world.category)
    rival = make_user(db, role=Role.ORGANISER)
    rival_event, rival_category = on_sale_event(db, rival, title="Rival Show")
    buyer = auth_header(client, make_user(db).username)
    confirmed_booking(client, buyer, rival_event, rival_category, quantity=3)
    rival_headers = auth_header(client, rival.username)

    mine = client.get("/api/v1/organiser/analytics/sales", headers=world.organiser_headers).json()
    assert [e["eventId"] for e in mine["events"]] == [str(world.event.id)] and mine["totalTicketsSold"] == 1
    theirs = client.get("/api/v1/organiser/analytics/sales", headers=rival_headers).json()
    assert [e["eventId"] for e in theirs["events"]] == [str(rival_event.id)] and theirs["totalTicketsSold"] == 3

    for path in ("sales", "revenue-velocity", "tier-demand"):
        peek = client.get(f"/api/v1/organiser/analytics/{path}?eventId={rival_event.id}",
                          headers=world.organiser_headers)
        assert peek.status_code == 404 and error_of(peek)["code"] == "EVENT_NOT_FOUND"
    own = client.get(f"/api/v1/organiser/analytics/sales?eventId={world.event.id}", headers=world.organiser_headers)
    assert own.status_code == 200 and own.json()["totalTicketsSold"] == 1


def test_analytics_are_for_organisers_only(client, world):
    for path in ("sales", "revenue-velocity", "tier-demand"):
        url = f"/api/v1/organiser/analytics/{path}"
        assert client.get(url).status_code == 401
        for headers in (world.headers, world.admin):
            response = client.get(url, headers=headers)
            assert response.status_code == 403 and error_of(response)["code"] == "FORBIDDEN"


def test_revenue_velocity_window_is_bounded(client, world):
    url = "/api/v1/organiser/analytics/revenue-velocity"
    assert client.get(f"{url}?days=6", headers=world.organiser_headers).status_code == 422
    assert client.get(f"{url}?days=91", headers=world.organiser_headers).status_code == 422
    assert client.get(f"{url}?days=90", headers=world.organiser_headers).status_code == 200


# ------------------------------------------------------------------ EaseVoice: MP3 recordings

MP3 = b"ID3\x04\x00\x00\x00\x00\x00\x00" + b"\xff\xfb\x90\x00" * 64  # tiny file with a valid MP3 signature


def put_audio(client, headers, page_key="home", data=MP3, filename=None):
    url = f"/api/v1/admin/narration/{page_key}/audio" + (f"?filename={filename}" if filename else "")
    return client.put(url, headers={**headers, "Content-Type": "audio/mpeg"}, content=data)


def with_narration(client, world, page_key="home", text="Welcome"):
    assert client.put(f"/api/v1/admin/narration/{page_key}", headers=world.admin,
                      json={"transcript": text}).status_code == 200


def test_admin_uploads_an_mp3_and_everyone_can_play_it(client, world):
    with_narration(client, world)
    assert client.get("/api/v1/narration/home").json()["hasAudio"] is False

    uploaded = put_audio(client, world.admin, filename="welcome.mp3")
    assert uploaded.status_code == 200, uploaded.text
    body = uploaded.json()
    assert body["hasAudio"] is True and body["audioFilename"] == "welcome.mp3"
    assert body["audioSizeBytes"] == len(MP3) and body["transcript"] == "Welcome"

    played = client.get("/api/v1/narration/home/audio")  # public, no login
    assert played.status_code == 200
    assert played.headers["content-type"] == "audio/mpeg" and played.content == MP3
    assert client.get("/api/v1/narration/home").json()["hasAudio"] is True

    # Editing the text keeps the recording; a new upload replaces it.
    client.put("/api/v1/admin/narration/home", headers=world.admin, json={"transcript": "New words"})
    assert client.get("/api/v1/narration/home").json()["hasAudio"] is True
    replacement = MP3 + b"\xff\xfb\x90\x00" * 10
    assert put_audio(client, world.admin, data=replacement).status_code == 200
    assert client.get("/api/v1/narration/home/audio").content == replacement


def test_mp3_upload_needs_existing_narration_text(client, world):
    response = put_audio(client, world.admin)
    assert response.status_code == 404 and error_of(response)["code"] == "NARRATION_NOT_FOUND"


def test_mp3_upload_is_validated(client, world):
    with_narration(client, world)
    for data, status, code in (
        (b"", 400, "AUDIO_EMPTY"),
        (b"RIFF....WAVEfmt " + b"\x00" * 64, 415, "AUDIO_NOT_MP3"),
        (b"%PDF-1.7 not audio at all", 415, "AUDIO_NOT_MP3"),
        (b"ID3" + b"\x00" * (5 * 1024 * 1024), 413, "AUDIO_TOO_LARGE"),
    ):
        response = put_audio(client, world.admin, data=data)
        assert response.status_code == status, (code, response.status_code, response.text[:200])
        assert error_of(response)["code"] == code
    assert client.get("/api/v1/narration/home/audio").status_code == 404  # nothing was stored


def test_mp3_file_name_is_cleaned(client, world):
    with_narration(client, world)
    uploaded = put_audio(client, world.admin, filename="..%2F..%2Fevil%22name.mp3")
    assert uploaded.status_code == 200
    name = uploaded.json()["audioFilename"]
    assert "/" not in name and '"' not in name and name.endswith(".mp3")
    assert put_audio(client, world.admin).json()["audioFilename"] == "home.mp3"  # default name


def test_only_admins_manage_recordings_but_anyone_can_listen(client, world):
    with_narration(client, world)
    assert put_audio(client, world.admin).status_code == 200
    for headers in (world.headers, world.organiser_headers):
        assert put_audio(client, headers).status_code == 403
        assert client.delete("/api/v1/admin/narration/home/audio", headers=headers).status_code == 403
    assert client.put("/api/v1/admin/narration/home/audio", content=MP3,
                      headers={"Content-Type": "audio/mpeg"}).status_code == 401
    assert client.delete("/api/v1/admin/narration/home/audio").status_code == 401
    assert client.get("/api/v1/narration/home/audio").status_code == 200  # still there, still public


def test_removing_a_recording_and_the_audit_trail(client, db, world):
    with_narration(client, world)
    assert put_audio(client, world.admin, filename="a.mp3").status_code == 200
    assert client.delete("/api/v1/admin/narration/home/audio", headers=world.admin).status_code == 204
    gone = client.get("/api/v1/narration/home/audio")
    assert gone.status_code == 404 and error_of(gone)["code"] == "NARRATION_AUDIO_NOT_FOUND"
    assert client.get("/api/v1/narration/home").json()["hasAudio"] is False
    again = client.delete("/api/v1/admin/narration/home/audio", headers=world.admin)
    assert again.status_code == 404 and error_of(again)["code"] == "NARRATION_AUDIO_NOT_FOUND"

    db.expire_all()
    rows = db.scalars(select(AuditLog).where(AuditLog.action.in_(
        ["NARRATION_AUDIO_UPLOADED", "NARRATION_AUDIO_REMOVED"])).order_by(AuditLog.created_at)).all()
    assert [r.action for r in rows] == ["NARRATION_AUDIO_UPLOADED", "NARRATION_AUDIO_REMOVED"]
    assert rows[0].before == {"hasAudio": False, "filename": None, "sizeBytes": None}
    assert rows[0].after["hasAudio"] is True and rows[0].after["sizeBytes"] == len(MP3)
    assert rows[1].before["hasAudio"] is True and rows[1].after["hasAudio"] is False
    assert all(r.actor_role == "ADMIN" for r in rows)
    assert all(MP3.hex() not in str(r.after) and MP3.hex() not in str(r.before) for r in rows)  # no raw bytes logged


# ------------------------------------------------------------------ Exports: EaseAnalytics (organiser)

def load_xlsx(response):
    from io import BytesIO

    from openpyxl import load_workbook
    return load_workbook(BytesIO(response.content))


def sheet_text(ws):
    return [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None]


def test_organiser_downloads_analytics_as_pdf(client, db, world):
    confirmed_booking(client, world.headers, world.event, world.category, quantity=2)
    response = client.get("/api/v1/organiser/analytics/export?format=pdf&days=7", headers=world.organiser_headers)
    assert response.status_code == 200, response.text[:300]
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith('attachment; filename="eventease-analytics-')
    assert response.headers["content-disposition"].endswith('.pdf"')
    assert response.content.startswith(b"%PDF") and len(response.content) > 2000
    default = client.get("/api/v1/organiser/analytics/export", headers=world.organiser_headers)  # pdf is the default
    assert default.headers["content-type"] == "application/pdf"


def test_organiser_downloads_analytics_as_excel_with_the_same_numbers(client, world):
    confirmed_booking(client, world.headers, world.event, world.category, quantity=2)
    response = client.get("/api/v1/organiser/analytics/export?format=xlsx&days=7", headers=world.organiser_headers)
    assert response.status_code == 200, response.text[:300]
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert response.headers["content-disposition"].endswith('.xlsx"')
    workbook = load_xlsx(response)
    assert workbook.sheetnames == ["Summary", "Sales per event", "Revenue trend", "Tier demand"]

    sales = workbook["Sales per event"]
    assert world.event.title in sheet_text(sales)
    row = [c.value for c in sales[2]]
    assert row[1:] == [1, 2, 100.0]  # 1 confirmed booking, 2 tickets, SGD 100.00
    trend = workbook["Revenue trend"]
    assert trend.max_row == 8  # header + 7 days
    assert trend.cell(row=8, column=3).value == 100.0 and trend.cell(row=8, column=4).value == 100.0
    assert len(workbook["Sales per event"]._charts) == 1 and len(trend._charts) == 2  # charts are included


def test_analytics_export_with_no_sales_still_works(client, world):
    for fmt in ("pdf", "xlsx"):
        response = client.get(f"/api/v1/organiser/analytics/export?format={fmt}", headers=world.organiser_headers)
        assert response.status_code == 200, (fmt, response.text[:200])


def test_analytics_export_can_be_limited_to_one_event_and_respects_ownership(client, db, world):
    other = make_user(db, role=Role.ORGANISER)
    rival_event, _ = on_sale_event(db, other)
    own = client.get(f"/api/v1/organiser/analytics/export?format=xlsx&eventId={world.event.id}",
                     headers=world.organiser_headers)
    assert own.status_code == 200
    assert f"Scope: {world.event.title}" in sheet_text(load_xlsx(own)["Summary"])
    rival = client.get(f"/api/v1/organiser/analytics/export?eventId={rival_event.id}",
                       headers=world.organiser_headers)
    assert rival.status_code == 404 and error_of(rival)["code"] == "EVENT_NOT_FOUND"


def test_analytics_export_is_for_organisers_only_and_validated(client, world):
    url = "/api/v1/organiser/analytics/export"
    assert client.get(url).status_code == 401
    for headers in (world.headers, world.admin):
        assert client.get(url, headers=headers).status_code == 403
    assert client.get(f"{url}?format=docx", headers=world.organiser_headers).status_code == 422
    assert client.get(f"{url}?days=3", headers=world.organiser_headers).status_code == 422


def test_analytics_downloads_are_audited(client, db, world):
    client.get("/api/v1/organiser/analytics/export?format=xlsx&days=14", headers=world.organiser_headers)
    db.expire_all()
    row = db.scalar(select(AuditLog).where(AuditLog.action == "ANALYTICS_EXPORTED"))
    assert row is not None and row.actor_role == "ORGANISER"
    assert row.after == {"format": "xlsx", "days": 14, "eventId": None}


# ------------------------------------------------------------------ Exports: EaseGuide audit (admin)

def lifecycle_request(client, world, needs="Wheelchair access"):
    booking_id = confirmed_booking(client, world.headers, world.event, world.category)
    request_id = request_assistance(client, world.headers, booking_id, supportNeeds=needs).json()["id"]
    url = f"/api/v1/admin/assistance-requests/{request_id}"
    client.patch(url, headers=world.admin, json={"assignedStaff": "Alex Tan"})
    client.patch(url, headers=world.admin, json={"status": "IN_PROGRESS"})
    return request_id


def test_admin_downloads_the_assistance_audit_report_as_pdf(client, world):
    lifecycle_request(client, world)
    response = client.get("/api/v1/admin/assistance-requests/export?format=pdf", headers=world.admin)
    assert response.status_code == 200, response.text[:300]
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith('attachment; filename="eventease-assistance-audit-')
    assert response.content.startswith(b"%PDF") and len(response.content) > 2000


def test_assistance_excel_has_every_request_and_its_status_history(client, world):
    request_id = lifecycle_request(client, world, needs="Needs step-free access")
    response = client.get("/api/v1/admin/assistance-requests/export?format=xlsx", headers=world.admin)
    assert response.status_code == 200, response.text[:300]
    workbook = load_xlsx(response)
    assert workbook.sheetnames == ["Summary", "Requests per day", "Request register", "Status history"]

    register = workbook["Request register"]
    values = [c.value for c in register[2]]
    assert values[0] == request_id and values[2] == world.customer.username
    assert values[3] == world.event.title and values[4] == "Needs step-free access"
    assert values[5:8] == ["Moderate", "In progress", "Alex Tan"]

    history = workbook["Status history"]
    steps = [(r[2].value, r[4].value or "", r[5].value or "") for r in history.iter_rows(min_row=2)]
    assert ("ASSISTANCE_REQUESTED", "", "Submitted") in steps
    assert ("ASSISTANCE_UPDATED", "Submitted", "Assigned") in steps
    assert ("ASSISTANCE_UPDATED", "Assigned", "In progress") in steps
    assert all(r[3].value in ("CUSTOMER", "ADMIN") for r in history.iter_rows(min_row=2))
    assert len(workbook["Summary"]._charts) == 2


def test_assistance_export_filters_by_status_and_date(client, world):
    lifecycle_request(client, world)  # in progress
    url = "/api/v1/admin/assistance-requests/export?format=xlsx"
    in_progress = load_xlsx(client.get(f"{url}&status=IN_PROGRESS", headers=world.admin))
    assert in_progress["Request register"].max_row == 2
    completed = load_xlsx(client.get(f"{url}&status=COMPLETED", headers=world.admin))
    assert completed["Request register"].max_row == 1  # header only, still a valid file

    today = utcnow().astimezone().date()
    inside = client.get(f"{url}&createdFrom={today - timedelta(days=1)}&createdTo={today + timedelta(days=1)}",
                        headers=world.admin)
    assert load_xlsx(inside)["Request register"].max_row == 2
    old = client.get(f"{url}&createdFrom=2000-01-01&createdTo=2000-01-31", headers=world.admin)
    assert load_xlsx(old)["Request register"].max_row == 1

    backwards = client.get(f"{url}&createdFrom=2026-02-01&createdTo=2026-01-01", headers=world.admin)
    assert backwards.status_code == 422 and error_of(backwards)["code"] == "INVALID_DATE_RANGE"
    assert client.get(f"{url}&status=NOPE", headers=world.admin).status_code == 422


def test_assistance_export_with_nothing_to_report_still_works(client, world):
    for fmt in ("pdf", "xlsx"):
        assert client.get(f"/api/v1/admin/assistance-requests/export?format={fmt}", headers=world.admin).status_code == 200


def test_assistance_export_is_admin_only(client, world):
    url = "/api/v1/admin/assistance-requests/export"
    assert client.get(url).status_code == 401
    for headers in (world.headers, world.organiser_headers):
        assert client.get(url, headers=headers).status_code == 403


def test_assistance_downloads_are_audited(client, db, world):
    lifecycle_request(client, world)
    client.get("/api/v1/admin/assistance-requests/export?format=pdf&status=IN_PROGRESS", headers=world.admin)
    db.expire_all()
    row = db.scalar(select(AuditLog).where(AuditLog.action == "ASSISTANCE_EXPORTED"))
    assert row is not None and row.actor_role == "ADMIN"
    assert row.after == {"format": "pdf", "status": "IN_PROGRESS", "createdFrom": None, "createdTo": None, "rows": 1}


def test_user_text_cannot_break_reports_or_run_as_excel_formulas(client, world):
    lifecycle_request(client, world, needs='=HYPERLINK("http://evil","click") <b>& more</b>')
    pdf = client.get("/api/v1/admin/assistance-requests/export?format=pdf", headers=world.admin)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")  # markup in user text didn't crash the PDF
    workbook = load_xlsx(client.get("/api/v1/admin/assistance-requests/export?format=xlsx", headers=world.admin))
    cell = workbook["Request register"].cell(row=2, column=5)
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")  # stored as text, never a formula


# ------------------------------------------------------------------ migration 0003

def test_migration_0003_is_reversible(alembic_config):
    schema = get_settings().DB_SCHEMA
    new_tables = {"assistance_requests", "narration_transcripts", "accessibility_settings"}

    command.downgrade(alembic_config, "0002")
    assert not new_tables & set(inspect(engine).get_table_names(schema=schema))

    command.upgrade(alembic_config, "head")
    assert new_tables <= set(inspect(engine).get_table_names(schema=schema))


def test_migration_0005_adds_and_removes_the_audio_columns(alembic_config):
    schema = get_settings().DB_SCHEMA
    audio = {"audio_data", "audio_filename", "audio_size_bytes"}

    command.downgrade(alembic_config, "0004")
    assert not audio & {c["name"] for c in inspect(engine).get_columns("narration_transcripts", schema=schema)}

    command.upgrade(alembic_config, "head")
    assert audio <= {c["name"] for c in inspect(engine).get_columns("narration_transcripts", schema=schema)}
