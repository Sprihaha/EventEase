"""Unit tests for the PDF/Excel report builders (no database needed)."""
import uuid
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services.reports import analytics_report, assistance_report
from app.services.reports import common as c

NOW = datetime(2026, 10, 19, 12, 0, tzinfo=timezone.utc)


def make_rows(count=12, needs="Wheelchair access", title="Jazz Night"):
    statuses = ["SUBMITTED", "ASSIGNED", "IN_PROGRESS", "COMPLETED", "CANCELLED"]
    rows, trail = [], []
    for i in range(count):
        created = NOW - timedelta(days=i % 4, hours=i)
        status = statuses[i % 5]
        request_id = uuid.uuid4()
        rows.append({
            "id": request_id, "booking_id": uuid.uuid4(), "status": status,
            "assistance_level": ["BASIC", "MODERATE", "FULL"][i % 3], "support_needs": needs,
            "assigned_staff": None if status == "SUBMITTED" else "Alex Tan", "created_at": created,
            "updated_at": created + timedelta(hours=2), "event": SimpleNamespace(title=title),
            "user_id": uuid.uuid4(), "requester_username": f"customer{i}",
        })
        trail.append({"at": created, "request_id": request_id, "action": "ASSISTANCE_REQUESTED",
                      "actor_role": "CUSTOMER", "from_status": None, "to_status": "SUBMITTED", "staff": None})
    return rows, trail


def analytics_inputs(empty=False):
    days = 7
    points = [{"day": date(2026, 10, 13) + timedelta(days=i), "tickets_sold": 0 if empty else i,
               "revenue_cents": 0 if empty else i * 5000,
               "cumulative_revenue_cents": 0 if empty else sum(range(i + 1)) * 5000} for i in range(days)]
    total = points[-1]["cumulative_revenue_cents"]
    events = [] if empty else [{"event_id": "e1", "title": "Jazz <Night> & Blues", "confirmed_bookings": 3,
                                "tickets_sold": 21, "revenue_cents": total}]
    tiers = [{"type": "STANDARD", "tickets_sold": 0 if empty else 21, "revenue_cents": total,
              "share_percent": 0.0 if empty else 100.0},
             {"type": "VIP", "tickets_sold": 0, "revenue_cents": 0, "share_percent": 0.0}]
    return dict(
        organiser="org1", scope="All my events", days=days, generated=NOW,
        sales={"currency": "SGD", "confirmed_bookings": 0 if empty else 3, "total_tickets_sold": 0 if empty else 21,
               "total_revenue_cents": total, "events": events},
        velocity={"currency": "SGD", "days": days, "total_revenue_cents": total,
                  "average_daily_revenue_cents": total // days, "points": points},
        tiers={"currency": "SGD", "total_tickets_sold": 0 if empty else 21, "tiers": tiers},
    )


# ---------------------------------------------------------------- helpers

def test_money_and_label_formatting():
    assert c.fmt_money(123456) == "SGD 1,234.56" and c.fmt_money(0) == "SGD 0.00"
    assert c.pretty("IN_PROGRESS") == "In progress" and c.pretty("VIP") == "Vip"
    assert c.fmt_dt(datetime(2026, 10, 19, 20, 0, tzinfo=timezone.utc)) == "20 Oct 2026 04:00"  # shown in Singapore time


def test_escaping_makes_markup_harmless():
    assert c.esc("<b>A & B</b>") == "&lt;b&gt;A &amp; B&lt;/b&gt;" and c.esc(None) == ""


def test_text_is_never_stored_as_an_excel_formula():
    from openpyxl import Workbook
    ws = Workbook().active
    for text in ("=1+1", "+SUM(A1)", "-2+3", '@cmd', "plain"):
        cell = c.put_text(ws, 1, 1, text)
        assert cell.data_type == "s" and cell.value == text


# ---------------------------------------------------------------- assistance summary

def test_summary_counts_and_timing():
    rows, _ = make_rows(10)
    summary = assistance_report.summarise(rows)
    assert summary["total"] == 10 and summary["open"] == 6
    assert summary["by_status"] == {"SUBMITTED": 2, "ASSIGNED": 2, "IN_PROGRESS": 2, "COMPLETED": 2, "CANCELLED": 2}
    assert summary["by_level"]["BASIC"] == 4 and sum(summary["by_level"].values()) == 10
    assert summary["avg_hours_to_complete"] == 2.0
    assert sum(n for _, n in summary["day_series"]) == 10
    days = [d for d, _ in summary["day_series"]]
    assert days == sorted(days) and (days[-1] - days[0]).days == len(days) - 1  # no gaps in the daily series


def test_summary_of_nothing():
    summary = assistance_report.summarise([])
    assert summary["total"] == 0 and summary["open"] == 0 and summary["day_series"] == []
    assert summary["avg_hours_to_complete"] is None
    assert assistance_report._fmt_hours(None) == "-" and assistance_report._fmt_hours(0.5) == "30 min"
    assert assistance_report._fmt_hours(5) == "5.0 h" and assistance_report._fmt_hours(72) == "3.0 days"


# ---------------------------------------------------------------- analytics files

def test_analytics_pdf_builds_with_and_without_data():
    for empty in (False, True):
        pdf = analytics_report.build_pdf(**analytics_inputs(empty))
        assert pdf.startswith(b"%PDF") and len(pdf) > 2000


def test_analytics_workbook_contents_and_charts():
    workbook = load_workbook(BytesIO(analytics_report.build_xlsx(**analytics_inputs())))
    assert workbook.sheetnames == ["Summary", "Sales per event", "Revenue trend", "Tier demand"]
    sales = workbook["Sales per event"]
    assert [cell.value for cell in sales[2]] == ["Jazz <Night> & Blues", 3, 21, 1050.0]
    assert sales.cell(row=3, column=4).value == 1050.0  # total row
    trend = workbook["Revenue trend"]
    assert trend.max_row == 8 and trend.cell(row=8, column=4).value == 1050.0
    assert len(sales._charts) == 1 and len(trend._charts) == 2 and len(workbook["Tier demand"]._charts) == 1


def test_analytics_workbook_with_no_sales_has_no_pointless_charts():
    workbook = load_workbook(BytesIO(analytics_report.build_xlsx(**analytics_inputs(empty=True))))
    assert len(workbook["Sales per event"]._charts) == 0 and len(workbook["Tier demand"]._charts) == 0


# ---------------------------------------------------------------- assistance files

def report_kwargs(rows, trail, total=None):
    return dict(requested_by="admin", filters="Status: all", rows=rows, total_matching=total or len(rows),
                trail=trail, generated=NOW)


def test_assistance_pdf_builds_with_and_without_data():
    rows, trail = make_rows(30)
    assert assistance_report.build_pdf(**report_kwargs(rows, trail)).startswith(b"%PDF")
    assert assistance_report.build_pdf(**report_kwargs([], [])).startswith(b"%PDF")


def test_assistance_pdf_survives_hostile_text_and_long_lists():
    rows, trail = make_rows(200, needs='<b>unclosed & "quoted" <script>x</script> ' + "word " * 80,
                            title="<para>Bad & title")
    pdf = assistance_report.build_pdf(**report_kwargs(rows, trail, total=5000))
    assert pdf.startswith(b"%PDF")


def test_assistance_workbook_has_all_rows_not_just_the_pdf_cap():
    rows, trail = make_rows(200)
    workbook = load_workbook(BytesIO(assistance_report.build_xlsx(**report_kwargs(rows, trail))))
    assert workbook["Request register"].max_row == 201 and workbook["Status history"].max_row == 201
    assert workbook["Request register"].auto_filter.ref and workbook["Request register"].freeze_panes == "A2"


def test_assistance_workbook_keeps_formulas_as_text():
    rows, trail = make_rows(2, needs='=HYPERLINK("http://evil","x")')
    workbook = load_workbook(BytesIO(assistance_report.build_xlsx(**report_kwargs(rows, trail))))
    cell = workbook["Request register"].cell(row=2, column=5)
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
