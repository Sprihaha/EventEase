"""Read an organiser's venue layout file (.xlsx or .csv) into sections, rows and seats.

The file lists seat ranges, one per line: Section, Row, From, To and an optional
Accessible column. This module only parses and validates; app.services.venues
stores the result.
"""
import csv
import io
from dataclasses import dataclass, field

from app.core.errors import AppError
from app.rules.venue import ROW_LABEL_MAX, SECTION_NAME_MAX

# The upload format. The template, the parser and the error messages all read these.
COLUMNS: dict[str, str] = {
    "Section": "Section name, e.g. the name printed on the venue plan",
    "Row": "Row label exactly as printed on the seats",
    "From": "First seat number in this row (whole number, 1 or more)",
    "To": "Last seat number in this row (whole number, same as From or larger)",
    "Accessible": "Optional. Wheelchair or easy-access seat numbers in this row, e.g. 1,2 or 5-8",
}
REQUIRED_COLUMNS = ("Section", "Row", "From", "To")
SUPPORTED_EXTENSIONS = (".xlsx", ".csv")

MAX_REPORTED_ERRORS = 50


@dataclass
class ParsedSeat:
    row: str
    number: int
    accessible: bool


@dataclass
class ParsedSection:
    name: str
    seats: list[ParsedSeat] = field(default_factory=list)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_rows(filename: str, content: bytes) -> list[list[str]]:
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise AppError(422, "UNREADABLE_FILE", "save the CSV file with UTF-8 encoding and upload it again")
        return [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text))]
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise AppError(422, "UNREADABLE_FILE", "the Excel file could not be opened; download the template and try again")
        try:
            return [[_cell(c) for c in row] for row in workbook.worksheets[0].iter_rows(values_only=True)]
        finally:
            workbook.close()
    raise AppError(422, "UNSUPPORTED_FILE", f"upload one of these file types: {', '.join(SUPPORTED_EXTENSIONS)}")


def _whole_number(text: str) -> int | None:
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else None


def _accessible_numbers(text: str) -> set[int] | None:
    """'1,2, 5-8' -> {1, 2, 5, 6, 7, 8}; None when the text is not in that form."""
    numbers: set[int] = set()
    for part in filter(None, (p.strip() for p in text.replace(";", ",").split(","))):
        low, _, high = part.partition("-")
        first, last = _whole_number(low), _whole_number(high or low)
        if first is None or last is None or first > last:
            return None
        numbers.update(range(first, last + 1))
    return numbers


def parse_layout(filename: str, content: bytes, max_seats: int) -> list[ParsedSection]:
    """Validate the whole file and report every bad line at once."""
    rows = _read_rows(filename, content)
    header = [c.casefold() for c in rows[0]] if rows else []
    missing = [name for name in REQUIRED_COLUMNS if name.casefold() not in header]
    if missing:
        raise AppError(422, "BAD_TEMPLATE",
                       f"the first row must contain the columns: {', '.join(COLUMNS)}",
                       {"errors": [f"missing column: {name}" for name in missing]})
    index = {name: header.index(name.casefold()) for name in COLUMNS if name.casefold() in header}

    def value(row: list[str], column: str) -> str:
        position = index.get(column)
        return row[position] if position is not None and position < len(row) else ""

    sections: dict[str, ParsedSection] = {}
    seen: dict[tuple[str, str, int], int] = {}
    errors: list[str] = []
    total = 0
    for line, row in enumerate(rows[1:], start=2):
        if not any(value(row, c) for c in COLUMNS):
            continue  # blank line
        section, row_label = value(row, "Section"), value(row, "Row")
        first, last = _whole_number(value(row, "From")), _whole_number(value(row, "To"))
        problems = []
        if not section:
            problems.append("Section is empty")
        elif len(section) > SECTION_NAME_MAX:
            problems.append(f"Section is longer than {SECTION_NAME_MAX} characters")
        if not row_label:
            problems.append("Row is empty")
        elif len(row_label) > ROW_LABEL_MAX:
            problems.append(f"Row is longer than {ROW_LABEL_MAX} characters")
        if first is None or last is None:
            problems.append("From and To must be whole numbers")
        elif first < 1 or last < first:
            problems.append("From must be 1 or more, and To must not be smaller than From")
        accessible = _accessible_numbers(value(row, "Accessible"))
        if accessible is None:
            problems.append("Accessible must list seat numbers like 1,2 or 5-8")
        elif first is not None and last is not None and accessible - set(range(first, last + 1)):
            problems.append("Accessible lists seat numbers outside From–To")
        if problems:
            errors.append(f"line {line}: " + "; ".join(problems))
            continue

        key_section = section.casefold()
        target = sections.setdefault(key_section, ParsedSection(name=section))
        for number in range(first, last + 1):
            key = (key_section, row_label.casefold(), number)
            if key in seen:
                errors.append(f"line {line}: seat {row_label}{number} in {section} is already listed on line {seen[key]}")
                break
            seen[key] = line
            target.seats.append(ParsedSeat(row=row_label, number=number, accessible=number in accessible))
        total = len(seen)
        if total > max_seats:
            errors.append(f"the layout has more than {max_seats} seats, the most one venue can have")
            break

    if errors:
        raise AppError(422, "INVALID_LAYOUT", "fix the lines listed below and upload the file again",
                       {"errors": errors[:MAX_REPORTED_ERRORS], "errorCount": len(errors)})
    if total == 0:
        raise AppError(422, "EMPTY_LAYOUT", "the file has no seats; add at least one row under the headings")
    return list(sections.values())


def template_workbook() -> bytes:
    """An empty upload template: the header row plus a sheet explaining each column."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    seats = workbook.active
    seats.title = "Seats"
    seats.append(list(COLUMNS))
    for cell in seats[1]:
        cell.font = Font(bold=True)
    for position, name in enumerate(COLUMNS, start=1):
        seats.column_dimensions[get_column_letter(position)].width = max(len(name) + 6, 14)
    seats.freeze_panes = "A2"

    guide = workbook.create_sheet("How to fill")
    guide.append(["Column", "What to enter", "Required"])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for name, description in COLUMNS.items():
        guide.append([name, description, "Yes" if name in REQUIRED_COLUMNS else "No"])
    guide.append([])
    guide.append(["Tip", "One line per row of seats. Lines with the same Section are grouped together."])
    guide.column_dimensions["A"].width = 14
    guide.column_dimensions["B"].width = 80

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
