"""EaseGuide audit export: PDF (charts + request register + status history) and Excel (same, with full data)."""
from collections import Counter
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, Spacer

from app.services.reports import common as c
from app.services.reports.analytics_report import _header, _widths

CONTENT_WIDTH = landscape(A4)[0] - 30 * mm
STATUS_ORDER = ["SUBMITTED", "ASSIGNED", "IN_PROGRESS", "COMPLETED", "CANCELLED"]
LEVEL_ORDER = ["BASIC", "MODERATE", "FULL"]
OPEN_STATUSES = {"SUBMITTED", "ASSIGNED", "IN_PROGRESS"}
PDF_REGISTER_LIMIT = 150
PDF_TRAIL_LIMIT = 150
DAY_CHART_LIMIT = 45
NOTE = ("Times are Singapore time (UTC+8). 'Time to complete' is measured from the request being submitted to its "
        "last update, for COMPLETED requests. The status history is taken from the system audit log; every change "
        "records the staff role that made it.")


def summarise(rows: list[dict]) -> dict:
    """Counts and timings used by both reports. Pure function, unit-tested."""
    statuses = Counter(c.enum_value(r["status"]) for r in rows)
    levels = Counter(c.enum_value(r["assistance_level"]) for r in rows)
    per_day = Counter(c.to_sgt(r["created_at"]).date() for r in rows)

    day_series: list[tuple[date, int]] = []
    if per_day:
        last = max(per_day)
        first = max(min(per_day), last - timedelta(days=DAY_CHART_LIMIT - 1))
        day_series = [(first + timedelta(days=i), per_day.get(first + timedelta(days=i), 0))
                      for i in range((last - first).days + 1)]

    hours = [(r["updated_at"] - r["created_at"]).total_seconds() / 3600
             for r in rows if c.enum_value(r["status"]) == "COMPLETED"]
    return {
        "total": len(rows),
        "by_status": {s: statuses.get(s, 0) for s in STATUS_ORDER},
        "by_level": {lv: levels.get(lv, 0) for lv in LEVEL_ORDER},
        "open": sum(n for s, n in statuses.items() if s in OPEN_STATUSES),
        "day_series": day_series,
        "avg_hours_to_complete": (sum(hours) / len(hours)) if hours else None,
    }


def _fmt_hours(hours: float | None) -> str:
    if hours is None:
        return "-"
    return f"{hours * 60:.0f} min" if hours < 1 else f"{hours:.1f} h" if hours < 48 else f"{hours / 24:.1f} days"


def _kpis(summary: dict) -> list[tuple[str, str]]:
    by = summary["by_status"]
    return [("Total requests", c.fmt_int(summary["total"])), ("Still open", c.fmt_int(summary["open"])),
            ("Completed", c.fmt_int(by["COMPLETED"])), ("Cancelled", c.fmt_int(by["CANCELLED"])),
            ("Avg time to complete", _fmt_hours(summary["avg_hours_to_complete"]))]


def _short_id(value) -> str:
    return str(value)[:8]


def build_pdf(*, requested_by: str, filters: str, rows: list[dict], total_matching: int, trail: list[dict],
              generated: datetime) -> bytes:
    st = c.styles()
    buffer = BytesIO()
    doc = c.new_doc(buffer, "EventEase assistance requests audit report", wide=True)
    summary = summarise(rows)
    third = CONTENT_WIDTH / 3 - 8
    by_status, by_level = summary["by_status"], summary["by_level"]

    story = [
        Paragraph("Assistance Requests - Audit Report", st["title"]),
        Paragraph(f"Requested by: {c.esc(requested_by)} (ADMIN) &nbsp;|&nbsp; Filters: {c.esc(filters)} &nbsp;|&nbsp; "
                  f"Generated {c.fmt_dt(generated)} SGT", st["sub"]),
        c.kpi_row(_kpis(summary), CONTENT_WIDTH),
        Spacer(1, 6),
    ]
    if total_matching > len(rows):
        story.append(Paragraph(f"Note: {c.fmt_int(total_matching)} requests match; this report covers the newest "
                               f"{c.fmt_int(len(rows))}. Narrow the filters to include the rest.", st["note"]))

    day_labels = [d.strftime("%d %b") for d, _ in summary["day_series"]]
    story.append(Paragraph("Requests at a glance", st["h2"]))
    story.append(c.Table([[
        c.chart_title(c.pie_chart([c.pretty(s) for s in STATUS_ORDER], [by_status[s] for s in STATUS_ORDER],
                                  width=third, height=150,
                                  slice_colours=[c.STATUS_COLOURS[s] for s in STATUS_ORDER if by_status[s] > 0]),
                      "By status"),
        c.chart_title(c.bar_chart([c.pretty(lv) for lv in LEVEL_ORDER], [by_level[lv] for lv in LEVEL_ORDER],
                                  width=third, height=150, colour=c.PALETTE[0]), "By assistance level"),
        c.chart_title(c.bar_chart(day_labels, [n for _, n in summary["day_series"]], width=third, height=150,
                                  colour=c.PALETTE[2]), "New requests per day"),
    ]], colWidths=[CONTENT_WIDTH / 3] * 3))

    status_rows = [["Status", "Requests", "Share"]]
    for s in STATUS_ORDER:
        share = 100 * by_status[s] / summary["total"] if summary["total"] else 0
        status_rows.append([c.pretty(s), c.fmt_int(by_status[s]), f"{share:.1f}%"])
    status_rows.append(["Total", c.fmt_int(summary["total"]), "100.0%" if summary["total"] else "0.0%"])
    story.append(KeepTogether([
        Paragraph("Status summary", st["h2"]),
        c.data_table(status_rows, [CONTENT_WIDTH * 0.25, CONTENT_WIDTH * 0.15, CONTENT_WIDTH * 0.15],
                     align_right={1, 2}, bold_last=True),
    ]))

    # ---- request register
    story.append(PageBreak())
    story.append(Paragraph("Request register", st["h2"]))
    register = [["Request", "Requester", "Event", "Needs", "Level", "Status", "Staff", "Submitted", "Last update"]]
    for r in rows[:PDF_REGISTER_LIMIT]:
        needs = r["support_needs"] if len(r["support_needs"]) <= 110 else r["support_needs"][:109] + "…"
        register.append([_short_id(r["id"]), Paragraph(c.esc(r["requester_username"]), st["cell"]),
                         Paragraph(c.esc(r["event"].title), st["cell"]), Paragraph(c.esc(needs), st["cell"]),
                         c.pretty(r["assistance_level"]), c.pretty(r["status"]),
                         Paragraph(c.esc(r["assigned_staff"] or "-"), st["cell"]),
                         c.fmt_dt(r["created_at"]), c.fmt_dt(r["updated_at"])])
    if len(register) == 1:
        register.append(["-", "No requests match these filters", "", "", "", "", "", "", ""])
    w = CONTENT_WIDTH
    story.append(c.data_table(register, [w * 0.07, w * 0.10, w * 0.14, w * 0.23, w * 0.07, w * 0.08, w * 0.09,
                                         w * 0.11, w * 0.11]))
    if len(rows) > PDF_REGISTER_LIMIT:
        story.append(Paragraph(f"Showing the newest {PDF_REGISTER_LIMIT} of {len(rows)}. "
                               "The Excel export contains every row.", st["note"]))

    # ---- status history from the audit log
    story.append(Paragraph("Status history (audit log)", st["h2"]))
    history = [["When", "Request", "Action", "By role", "From", "To", "Staff"]]
    for t in trail[:PDF_TRAIL_LIMIT]:
        history.append([c.fmt_dt(t["at"]), _short_id(t["request_id"]), t["action"].replace("ASSISTANCE_", "").title(),
                        t["actor_role"] or "-", c.pretty(t["from_status"]) if t["from_status"] else "-",
                        c.pretty(t["to_status"]) if t["to_status"] else "-",
                        Paragraph(c.esc(t["staff"] or "-"), st["cell"])])
    if len(history) == 1:
        history.append(["-", "No recorded changes", "", "", "", "", ""])
    story.append(c.data_table(history, [w * 0.15, w * 0.10, w * 0.14, w * 0.12, w * 0.14, w * 0.14, w * 0.21]))
    if len(trail) > PDF_TRAIL_LIMIT:
        story.append(Paragraph(f"Showing the newest {PDF_TRAIL_LIMIT} of {len(trail)} changes. "
                               "The Excel export contains the full history.", st["note"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(NOTE, st["note"]))

    footer = c.page_furniture("Assistance Requests Audit Report", generated)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- Excel

def build_xlsx(*, requested_by: str, filters: str, rows: list[dict], total_matching: int, trail: list[dict],
               generated: datetime) -> bytes:
    summary = summarise(rows)
    wb = Workbook()

    # ---- Summary + charts
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "EventEase - Assistance Requests Audit Report"
    ws["A1"].font = Font(bold=True, size=16)
    c.put_text(ws, 2, 1, f"Requested by: {requested_by} (ADMIN)")
    c.put_text(ws, 3, 1, f"Filters: {filters}")
    ws["A4"] = f"Generated: {c.fmt_dt(generated)} SGT"
    if total_matching > len(rows):
        ws["A5"] = f"Note: {total_matching} requests match; this workbook covers the newest {len(rows)}."
    _header(ws, 7, ["Measure", "Value"])
    for i, (label, value) in enumerate(_kpis(summary), start=8):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)

    _header(ws, 15, ["Status", "Requests"])
    for i, s in enumerate(STATUS_ORDER, start=16):
        ws.cell(row=i, column=1, value=c.pretty(s))
        ws.cell(row=i, column=2, value=summary["by_status"][s])
    _header(ws, 15 + len(STATUS_ORDER) + 2, ["Assistance level", "Requests"])
    level_start = 15 + len(STATUS_ORDER) + 3
    for i, lv in enumerate(LEVEL_ORDER, start=level_start):
        ws.cell(row=i, column=1, value=c.pretty(lv))
        ws.cell(row=i, column=2, value=summary["by_level"][lv])
    _widths(ws, [30, 18])

    if summary["total"]:
        pie = PieChart()
        pie.title = "Requests by status"
        pie.add_data(Reference(ws, min_col=2, min_row=15, max_row=15 + len(STATUS_ORDER)), titles_from_data=True)
        pie.set_categories(Reference(ws, min_col=1, min_row=16, max_row=15 + len(STATUS_ORDER)))
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.dataLabels.showVal = pie.dataLabels.showCatName = pie.dataLabels.showSerName = False
        pie.dataLabels.showLegendKey = False
        pie.height, pie.width = 8, 12
        ws.add_chart(pie, "D2")
        bars = BarChart()
        bars.type, bars.title, bars.legend = "col", "Requests by assistance level", None
        bars.add_data(Reference(ws, min_col=2, min_row=level_start - 1, max_row=level_start + len(LEVEL_ORDER) - 1),
                      titles_from_data=True)
        bars.set_categories(Reference(ws, min_col=1, min_row=level_start,
                                      max_row=level_start + len(LEVEL_ORDER) - 1))
        bars.height, bars.width = 8, 12
        ws.add_chart(bars, "D19")

    # ---- Per-day
    wd = wb.create_sheet("Requests per day")
    _header(wd, 1, ["Day", "New requests"])
    for i, (day, n) in enumerate(summary["day_series"], start=2):
        wd.cell(row=i, column=1, value=day).number_format = "dd mmm yyyy"
        wd.cell(row=i, column=2, value=n)
    _widths(wd, [16, 14])
    if summary["day_series"]:
        chart = BarChart()
        chart.type, chart.title, chart.legend = "col", "New requests per day", None
        chart.add_data(Reference(wd, min_col=2, min_row=1, max_row=len(summary["day_series"]) + 1),
                       titles_from_data=True)
        chart.set_categories(Reference(wd, min_col=1, min_row=2, max_row=len(summary["day_series"]) + 1))
        chart.height, chart.width = 8, 18
        wd.add_chart(chart, "D2")

    # ---- Register (all rows)
    wr = wb.create_sheet("Request register")
    heads = ["Request ID", "Booking ID", "Requester", "Event", "Support needs", "Level", "Status", "Assigned staff",
             "Submitted (SGT)", "Last update (SGT)"]
    _header(wr, 1, heads)
    for i, r in enumerate(rows, start=2):
        c.put_text(wr, i, 1, str(r["id"]))
        c.put_text(wr, i, 2, str(r["booking_id"]))
        c.put_text(wr, i, 3, r["requester_username"])
        c.put_text(wr, i, 4, r["event"].title)
        cell = c.put_text(wr, i, 5, r["support_needs"])
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        c.put_text(wr, i, 6, c.pretty(r["assistance_level"]))
        c.put_text(wr, i, 7, c.pretty(r["status"]))
        c.put_text(wr, i, 8, r["assigned_staff"] or "")
        c.put_text(wr, i, 9, c.fmt_dt(r["created_at"]))
        c.put_text(wr, i, 10, c.fmt_dt(r["updated_at"]))
    wr.freeze_panes = "A2"
    wr.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{max(len(rows) + 1, 2)}"
    _widths(wr, [38, 38, 18, 28, 50, 11, 13, 18, 18, 18])

    # ---- Audit trail (all rows)
    wt = wb.create_sheet("Status history")
    theads = ["When (SGT)", "Request ID", "Action", "By role", "From status", "To status", "Assigned staff"]
    _header(wt, 1, theads)
    for i, t in enumerate(trail, start=2):
        c.put_text(wt, i, 1, c.fmt_dt(t["at"]))
        c.put_text(wt, i, 2, str(t["request_id"]))
        c.put_text(wt, i, 3, t["action"])
        c.put_text(wt, i, 4, t["actor_role"] or "")
        c.put_text(wt, i, 5, c.pretty(t["from_status"]) if t["from_status"] else "")
        c.put_text(wt, i, 6, c.pretty(t["to_status"]) if t["to_status"] else "")
        c.put_text(wt, i, 7, t["staff"] or "")
    wt.freeze_panes = "A2"
    wt.auto_filter.ref = f"A1:{get_column_letter(len(theads))}{max(len(trail) + 1, 2)}"
    _widths(wt, [18, 38, 24, 12, 14, 14, 20])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
